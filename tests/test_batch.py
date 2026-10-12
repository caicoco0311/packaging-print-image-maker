from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("prepare_batch_tasks", ROOT / "scripts" / "prepare_batch_tasks.py")
prepare_batch_tasks = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(prepare_batch_tasks)


def load_script(name: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


workbook_to_batch_manifest = load_script("workbook_to_batch_manifest")
update_batch_results = load_script("update_batch_results")


def manifest() -> dict:
    return {
        "schema_version": "1.0",
        "batch_id": "batch-test",
        "source_workbook": "input.xlsx",
        "jobs": [
            {"row_number": 2, "task_id": "WINE-001", "raw_input": "红酒 + Label", "status": "pending", "overrides": {}},
            {"row_number": 3, "task_id": "SKIP-001", "raw_input": "咖啡 + Box", "status": "skipped", "overrides": {}},
            {"row_number": 4, "task_id": "BAD-001", "raw_input": "", "status": "blocked", "error": "missing user_input", "overrides": {}},
        ],
    }


class BatchPrepareTests(unittest.TestCase):
    def test_research_key_groups_product_and_separates_category(self):
        self.assertEqual(prepare_batch_tasks.research_key("red wine + Label"), prepare_batch_tasks.research_key("red wine 标签"))
        self.assertNotEqual(prepare_batch_tasks.research_key("red wine + Label"), prepare_batch_tasks.research_key("red wine + Box"))

    def test_xlsx_unicode_manifest_and_results_workbook(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = ROOT / "assets" / "batch-input-template.xlsx"
            manifest_path = workbook_to_batch_manifest.convert(source, root / "manifest.json", batch_id="batch-test")
            data = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(data["jobs"][0]["raw_input"], "红酒 + Label")
            self.assertEqual(data["jobs"][1]["overrides"]["material"], "cotton paper")
            self.assertEqual(data["jobs"][2]["status"], "skipped")
            state_path = prepare_batch_tasks.prepare(manifest_path, root / "run")
            result = update_batch_results.update(source, state_path, root / "results.xlsx")
            workbook = load_workbook(result, read_only=True)
            self.assertIn("Batch Results", workbook.sheetnames)
            self.assertEqual(workbook["Batch Results"]["A2"].value, "WINE-001")
            workbook.close()

    def test_isolation_and_blocked_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(manifest(), ensure_ascii=False), encoding="utf-8")
            state_path = prepare_batch_tasks.prepare(manifest_path, root / "out")
            state = json.loads(state_path.read_text(encoding="utf-8"))
            self.assertEqual(state["jobs"]["WINE-001"]["status"], "pending")
            self.assertEqual(state["jobs"]["SKIP-001"]["status"], "skipped")
            self.assertEqual(state["jobs"]["BAD-001"]["status"], "blocked")
            self.assertTrue((root / "out" / "WINE-001" / "row_input.json").is_file())
            self.assertFalse((root / "out" / "SKIP-001").exists())

    def test_resume_unchanged_and_reset_changed_row(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest_path = root / "manifest.json"
            data = manifest()
            manifest_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            state_path = prepare_batch_tasks.prepare(manifest_path, root / "out")
            state = json.loads(state_path.read_text(encoding="utf-8"))
            state["jobs"]["WINE-001"].update({"status": "complete", "stage": "validated", "images": ["1.png"]})
            state_path.write_text(json.dumps(state), encoding="utf-8")
            resumed = json.loads(prepare_batch_tasks.prepare(manifest_path, root / "out").read_text(encoding="utf-8"))
            self.assertEqual(resumed["jobs"]["WINE-001"]["status"], "complete")
            data["jobs"][0]["raw_input"] = "红酒 + 棉纸 Label"
            manifest_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            changed = json.loads(prepare_batch_tasks.prepare(manifest_path, root / "out").read_text(encoding="utf-8"))
            self.assertEqual(changed["jobs"]["WINE-001"]["status"], "pending")
            self.assertTrue(changed["jobs"]["WINE-001"]["input_changed"])
            self.assertEqual(changed["jobs"]["WINE-001"]["images"], [])

    def test_unsafe_task_id_is_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            data = manifest()
            data["jobs"] = [{"row_number": 2, "task_id": "../escape", "raw_input": "红酒 + Label", "status": "pending", "overrides": {}}]
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            state = json.loads(prepare_batch_tasks.prepare(manifest_path, root / "out").read_text(encoding="utf-8"))
            self.assertEqual(state["jobs"]["../escape"]["status"], "blocked")
            self.assertFalse((root / "escape").exists())


if __name__ == "__main__":
    unittest.main()
