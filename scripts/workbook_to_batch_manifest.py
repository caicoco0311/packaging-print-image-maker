#!/usr/bin/env python3
"""Convert an XLSX, CSV, or TSV batch sheet into an isolated-task manifest."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


HEADERS = ("task_id", "user_input")
OVERRIDES = ("material", "palette", "finish", "notes", "output_name")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _enabled(value) -> tuple[bool, bool]:
    if value is None or _text(value) == "":
        return True, True
    if isinstance(value, bool):
        return value, True
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value), True
    normalized = _text(value).lower()
    if normalized in {"true", "yes", "y", "1"}:
        return True, True
    if normalized in {"false", "no", "n", "0"}:
        return False, True
    return False, False


def _rows(path: Path, sheet_name: str | None) -> tuple[str, list[list]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        return sheet.title, rows
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return path.stem, [list(row) for row in csv.reader(stream, delimiter=delimiter)]
    if suffix == ".xls":
        raise ValueError("legacy .xls is unsupported; save as .xlsx, .csv, or .tsv")
    raise ValueError("batch input must be .xlsx, .csv, or .tsv")


def convert(input_path: Path, output_path: Path, sheet_name: str | None = None, batch_id: str | None = None) -> Path:
    source_sheet, rows = _rows(input_path, sheet_name)
    if not rows:
        raise ValueError("input worksheet is empty")
    headers = [_text(value).lower() for value in rows[0]]
    index = {name: position for position, name in enumerate(headers) if name}
    missing = [name for name in HEADERS if name not in index]
    if missing:
        raise ValueError(f"missing required columns: {', '.join(missing)}")

    raw_jobs = []
    for row_number, row in enumerate(rows[1:], 2):
        if not any(_text(value) for value in row):
            continue
        value = lambda name: row[index[name]] if name in index and index[name] < len(row) else None
        enabled, enabled_valid = _enabled(value("enabled"))
        raw_jobs.append({
            "row_number": row_number,
            "task_id": _text(value("task_id")),
            "raw_input": _text(value("user_input")),
            "enabled": enabled,
            "enabled_valid": enabled_valid,
            "overrides": {name: _text(value(name)) for name in OVERRIDES},
        })

    counts: dict[str, int] = {}
    for job in raw_jobs:
        if job["task_id"]:
            counts[job["task_id"]] = counts.get(job["task_id"], 0) + 1
    jobs = []
    for job in raw_jobs:
        errors = []
        if not job["task_id"]:
            errors.append("missing task_id")
        if not job["raw_input"]:
            errors.append("missing user_input")
        if not job.pop("enabled_valid"):
            errors.append("invalid enabled value")
        if job["task_id"] and counts[job["task_id"]] > 1:
            errors.append("duplicate task_id")
        job["status"] = "blocked" if errors else "pending" if job["enabled"] else "skipped"
        job["error"] = "; ".join(errors)
        jobs.append(job)

    manifest = {
        "schema_version": "1.0",
        "batch_id": batch_id or datetime.now(timezone.utc).strftime("batch-%Y%m%d%H%M%S"),
        "source_workbook": str(input_path.resolve()),
        "source_sheet": source_sheet,
        "jobs": jobs,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--batch-id")
    args = parser.parse_args()
    print(convert(args.input, args.output, args.sheet, args.batch_id))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
