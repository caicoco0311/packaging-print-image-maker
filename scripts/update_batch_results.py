#!/usr/bin/env python3
"""Write resumable batch state into a new XLSX results workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill


HEADERS = ("task_id", "row_number", "status", "stage", "research_key", "reference_mode", "product_title", "detail_page_keywords", "physical_qa", "retries", "output_dir", "error", *(f"image_{i}" for i in range(1, 7)))


def update(input_path: Path, state_path: Path, output_path: Path) -> Path:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("output workbook must not overwrite the input workbook")
    state = json.loads(state_path.read_text(encoding="utf-8"))
    workbook = load_workbook(input_path) if input_path.suffix.lower() == ".xlsx" else Workbook()
    if "Batch Results" in workbook.sheetnames:
        del workbook["Batch Results"]
    sheet = workbook.create_sheet("Batch Results")
    sheet.sheet_view.showGridLines = False
    sheet.append(HEADERS)
    for task_id, job in state.get("jobs", {}).items():
        images = list(job.get("images", []))[:6]
        identity_path = Path(job.get("output_dir", "")) / "identity.json"
        identity = json.loads(identity_path.read_text(encoding="utf-8")) if identity_path.is_file() else {}
        sheet.append((
            task_id, job.get("row_number", ""), job.get("status", ""), job.get("stage", ""),
            job.get("research_key", ""), identity.get("reference_mode", job.get("reference_mode", "")),
            identity.get("normalized_product_title", job.get("normalized_product_title", "")),
            ", ".join(identity.get("detail_page_keywords", job.get("detail_page_keywords", []))),
            job.get("physical_qa", ""), json.dumps(job.get("retries", {}), ensure_ascii=False), job.get("output_dir", ""), job.get("error", ""),
            *(images + [""] * (6 - len(images))),
        ))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    widths = (18, 12, 14, 18, 24, 16, 38, 50, 28, 18, 36, 34, 28, 28, 28, 28, 28, 28)
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    last = max(sheet.max_row, 2)
    sheet.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=['"complete"'], fill=PatternFill("solid", fgColor="E2F0D9")))
    sheet.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=['"blocked"'], fill=PatternFill("solid", fgColor="FCE4D6")))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--state", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    print(update(args.input, args.state, args.output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
